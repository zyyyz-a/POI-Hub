from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "POI-Hub-商家成交话术与零基础操作手册.docx"
SOURCES = [
    ROOT / "docs" / "merchant-sales-playbook.md",
    ROOT / "docs" / "novice-operator-sop.md",
    ROOT / "docs" / "pilot-store-checklist-disifengshang.md",
]

FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17365D"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
PALE_RED = "FDECEC"
RED = "9B1C1C"
BORDER = "CBD5E1"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, *, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, italic: bool | None = None,
                 latin: str = FONT_LATIN, cjk: str = FONT_CJK) -> None:
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_style(style, *, size: float, color: str = "000000", bold: bool = False,
                    before: float = 0, after: float = 6, line: float = 1.25,
                    keep_with_next: bool = False) -> None:
    style.font.name = FONT_LATIN
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_with_next


def configure_styles(doc: Document) -> None:
    configure_style(doc.styles["Normal"], size=11, after=6, line=1.25)
    configure_style(doc.styles["Title"], size=30, color=NAVY, bold=True, after=8, line=1.0,
                    keep_with_next=True)
    doc.styles["Title"].paragraph_format.page_break_before = False
    configure_style(doc.styles["Subtitle"], size=13.5, color=MUTED, after=18, line=1.15,
                    keep_with_next=True)
    configure_style(doc.styles["Heading 1"], size=16, color=BLUE, bold=True, before=18,
                    after=10, line=1.15, keep_with_next=True)
    configure_style(doc.styles["Heading 2"], size=13, color=BLUE, bold=True, before=14,
                    after=7, line=1.15, keep_with_next=True)
    configure_style(doc.styles["Heading 3"], size=12, color=DARK_BLUE, bold=True, before=10,
                    after=5, line=1.15, keep_with_next=True)
    configure_style(doc.styles["Heading 4"], size=11, color=DARK_BLUE, bold=True, before=8,
                    after=4, line=1.15, keep_with_next=True)

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        configure_style(style, size=11, after=4, line=1.25)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.tab_stops.add_tab_stop(Inches(0.375))

    if "Checklist" not in doc.styles:
        checklist = doc.styles.add_style("Checklist", WD_STYLE_TYPE.PARAGRAPH)
    else:
        checklist = doc.styles["Checklist"]
    configure_style(checklist, size=11, after=4, line=1.25)
    checklist.paragraph_format.left_indent = Inches(0.22)
    checklist.paragraph_format.first_line_indent = Inches(-0.22)

    if "Code Block" not in doc.styles:
        code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = doc.styles["Code Block"]
    configure_style(code, size=9.5, color="243447", after=0, line=1.10)
    code.font.name = "Consolas"
    code._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
    code._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
    code._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(1)
    code.paragraph_format.space_after = Pt(1)


def create_decimal_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    lvl.extend([start, num_fmt, lvl_text, jc, p_pr])
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return num_id


def add_numbered_paragraph(doc: Document, text: str, num_id: int):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p_pr.append(num_pr)
    add_inline(paragraph, text)
    return paragraph


def create_bullet_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    lvl.extend([start, num_fmt, lvl_text, jc, p_pr])
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def add_bullet_paragraph(doc: Document, text: str, num_id: int):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p_pr.append(num_pr)
    add_inline(paragraph, text)
    return paragraph


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = BORDER, size: int = 6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def choose_widths(column_count: int) -> list[int]:
    patterns = {
        2: [2700, 6660],
        3: [2100, 4200, 3060],
        4: [1900, 3100, 1600, 2760],
    }
    if column_count in patterns:
        return patterns[column_count]
    base = TABLE_WIDTH_DXA // column_count
    widths = [base] * column_count
    widths[-1] += TABLE_WIDTH_DXA - sum(widths)
    return widths


def apply_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError("Table widths must total 9360 DXA")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = tr_pr.find(qn("w:cantSplit"))
        if cant_split is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_left_border(paragraph, color: str, size: int = 18, space: int = 8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)
    p_bdr.append(left)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_LATIN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    fonts.set(qn("w:eastAsia"), FONT_CJK)
    r_pr.extend([fonts, color, underline])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, *, base_bold: bool = False) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, bold=base_bold)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=10, color=DARK_BLUE, latin="Consolas")
            set_paragraph_shading(paragraph, "F6F8FA")
        else:
            found = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if found:
                add_hyperlink(paragraph, found.group(1), found.group(2))
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, bold=base_bold)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    normalized = [row + [""] * (columns - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=columns)
    apply_table_geometry(table, choose_widths(columns))
    set_table_borders(table)
    for row_index, values in enumerate(normalized):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.10
            add_inline(p, value, base_bold=row_index == 0)
            for run in p.runs:
                set_run_font(run, size=9.4 if columns >= 4 else 9.8, bold=(row_index == 0))
            if row_index == 0:
                p.paragraph_format.keep_with_next = True
                set_cell_shading(cell, LIGHT_BLUE)
        if row_index == 0:
            repeat_table_header(table.rows[0])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def parse_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip().strip("|")
        cells = [cell.strip() for cell in raw.split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def add_code_block(doc: Document, code_lines: list[str]) -> None:
    for line in code_lines:
        p = doc.add_paragraph(style="Code Block")
        set_paragraph_shading(p, LIGHT_GRAY)
        run = p.add_run(line or " ")
        set_run_font(run, size=9.3, color="243447", latin="Consolas")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_blockquote(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    set_paragraph_shading(p, "EEF4FA")
    set_left_border(p, BLUE)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=DARK_BLUE)


def add_markdown(doc: Document, source: Path, *, part_number: int) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines: list[str] = []
    first_h1 = True
    current_decimal_num_id: int | None = None
    bullet_num_id = create_bullet_numbering(doc)
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            current_decimal_num_id = None
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            current_decimal_num_id = None
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2)
            if level == 1 and first_h1:
                p = doc.add_paragraph()
                p.paragraph_format.page_break_before = True
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(5)
                run = p.add_run(f"第 {part_number} 部分")
                set_run_font(run, size=10.5, bold=True, color=BLUE)
                p = doc.add_paragraph(style="Title")
                add_inline(p, title)
                if title.startswith("缔丝风尚潮色染烫"):
                    for run in p.runs:
                        set_run_font(run, size=22, bold=True, color=NAVY)
                elif len(title) > 16:
                    for run in p.runs:
                        set_run_font(run, size=24, bold=True, color=NAVY)
                first_h1 = False
            else:
                mapped = min(max(level - 1, 1), 4)
                p = doc.add_paragraph(style=f"Heading {mapped}")
                add_inline(p, title)
            current_decimal_num_id = None
            index += 1
            continue
        if stripped.startswith(">"):
            add_blockquote(doc, stripped[1:].strip())
            current_decimal_num_id = None
            index += 1
            continue
        check = re.match(r"^- \[([ xX])\]\s+(.+)$", stripped)
        if check:
            p = doc.add_paragraph(style="Checklist")
            marker = "☒" if check.group(1).lower() == "x" else "☐"
            run = p.add_run(f"{marker} ")
            set_run_font(run, size=11, color=BLUE)
            add_inline(p, check.group(2))
            current_decimal_num_id = None
            index += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            add_bullet_paragraph(doc, bullet.group(1), bullet_num_id)
            current_decimal_num_id = None
            index += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            if current_decimal_num_id is None:
                current_decimal_num_id = create_decimal_numbering(doc)
            add_numbered_paragraph(doc, numbered.group(1), current_decimal_num_id)
            index += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, stripped)
        current_decimal_num_id = None
        index += 1


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run_el = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MUTED)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    rpr.extend([color, size])
    run_el.append(rpr)
    text = OxmlElement("w:t")
    text.text = "1"
    run_el.append(text)
    fld.append(run_el)
    paragraph._p.append(fld)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color=MUTED)


def set_running_furniture(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("POI Hub｜商家成交与运营交付手册")
    set_run_font(run, size=9, color=MUTED)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), "D7DBE2")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    add_page_number(p)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("POI HUB · 单店受控试点")
    set_run_font(run, size=10.5, bold=True, color=BLUE)

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(8)
    add_inline(p, "商家成交话术与零基础操作手册")

    p = doc.add_paragraph(style="Subtitle")
    add_inline(p, "从商家沟通、微信开店、腾讯地图/服务 POI，到商品、核销、对账和异常升级")

    table = doc.add_table(rows=2, cols=2)
    values = [
        ("适用人员", "销售、零基础操作员、门店店员、平台负责人"),
        ("版本与日期", "V1.0 · 2026-08-25"),
    ]
    for row_index, (label, value) in enumerate(values):
        for column_index, text in enumerate((label, value)):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            set_run_font(run, size=10.5, bold=(column_index == 0), color=(DARK_BLUE if column_index == 0 else "000000"))
            set_cell_shading(cell, LIGHT_BLUE if column_index == 0 else WHITE)
    apply_table_geometry(table, [2300, 7060])
    set_table_borders(table)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    set_paragraph_shading(p, PALE_RED)
    set_left_border(p, RED)
    run = p.add_run("重要边界：本手册适用于一家店的人工受控试点。消费者货款由微信小店按规则结算给商家；POI Hub 服务费另行收取。真实微信授权、敏感凭据和永久性错误必须由负责人/技术人员处理。")
    set_run_font(run, size=11, bold=True, color=RED)

    p = doc.add_paragraph(style="Heading 1")
    add_inline(p, "使用顺序")
    cover_num_id = create_decimal_numbering(doc)
    for text in (
        "先让销售使用第 1 部分，只说真实能力和边界。",
        "签约后让操作员从第 2 部分第 0 节开始逐关执行。",
        "缔丝风尚潮色染烫试点直接使用第 3 部分逐项打勾。",
        "任何步骤出现红线或未知字段，立即停手并按异常格式升级。",
    ):
        add_numbered_paragraph(doc, text, cover_num_id)

    p = doc.add_paragraph(style="Heading 1")
    add_inline(p, "内容导航")
    cover_bullet_num_id = create_bullet_numbering(doc)
    for text in (
        "第 1 部分：商家沟通、异议处理、报价与销售交接",
        "第 2 部分：零基础操作员端到端 SOP",
        "第 3 部分：缔丝风尚潮色染烫单店执行卡",
    ):
        add_bullet_paragraph(doc, text, cover_bullet_num_id)


def add_part_break(doc: Document) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build() -> Path:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    set_running_furniture(section)
    configure_styles(doc)
    doc.core_properties.title = "POI Hub 商家成交话术与零基础操作手册"
    doc.core_properties.subject = "微信小店本地生活、腾讯地图/服务 POI、商品、核销与对账单店试点"
    doc.core_properties.author = "POI Hub"
    doc.core_properties.keywords = "微信小店, 本地生活, POI, 商家话术, 操作SOP"

    add_cover(doc)
    for index, source in enumerate(SOURCES, start=1):
        add_markdown(doc, source, part_number=index)

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
